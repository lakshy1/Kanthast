"""Build Kanthast_Class_<N>.xlsx from the per-chapter script JSON.

Layout: one workbook per class, one sheet per subject.
Source:  source/class-<N>/<subject-slug>/chNN.json  (see source/class-10/science/BRIEF.md)
Usage:   py -3.12 build_workbook.py 10
"""
import json
import math
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

WORDS_PER_MIN = 130      # narration pace for Class IX-X explainers
SCENE_BEAT_SECONDS = 2   # visual-only pause per scene transition
ROMAN = {1: "I", 2: "II", 3: "III", 4: "IV", 5: "V", 6: "VI", 7: "VII", 8: "VIII", 9: "IX", 10: "X"}
# Sheet order follows the app's subject order in Frontend/src/utils/schoolTrack.js.
SUBJECT_ORDER = ["mathematics", "science", "english", "social-science", "hindi", "sanskrit",
                 "information-technology", "evs", "social-studies", "general-knowledge", "computer-basics"]

ROOT = Path(__file__).parent


def fmt(seconds):
    return f"{seconds // 60}:{seconds % 60:02d}"


def scene_seconds(scene):
    words = sum(len(line["text"].split()) for line in scene["lines"])
    return math.ceil(words / WORDS_PER_MIN * 60) + SCENE_BEAT_SECONDS


def render_script(topic):
    parts, clock = [], 0
    for i, scene in enumerate(topic["scenes"], 1):
        secs = scene_seconds(scene)
        parts.append(f"SCENE {i}  ({fmt(clock)}–{fmt(clock + secs)})")
        parts.append(f"[ANIMATION] {scene['animation'].strip()}")
        parts.extend(f"{line['speaker'].upper()}: \"{line['text'].strip()}\"" for line in scene["lines"])
        parts.append("")
        clock += secs
    # Round the total up to the next 5 s so the Duration column reads cleanly.
    return "\n".join(parts).rstrip(), int(math.ceil(clock / 5) * 5)


def load_subject(folder):
    topics = []
    for path in sorted(folder.glob("ch*.json")):
        topics.extend(json.loads(path.read_text(encoding="utf-8")))
    return sorted(topics, key=lambda t: (t["chapter_no"], t["topic_no"]))


def build(class_no):
    source = ROOT / "source" / f"class-{class_no}"
    folders = sorted((p for p in source.iterdir() if p.is_dir()),
                     key=lambda p: SUBJECT_ORDER.index(p.name) if p.name in SUBJECT_ORDER else 99)
    wb = Workbook()
    wb.remove(wb.active)

    head_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    head_fill = PatternFill("solid", start_color="0E7490")
    top_wrap = Alignment(wrap_text=True, vertical="top")
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for folder in folders:
        topics = load_subject(folder)
        if not topics:
            continue
        subject = folder.name.replace("-", " ").title()
        # Arial has no Devanagari glyphs; Nirmala UI ships with Windows and covers both scripts.
        font_name = "Nirmala UI" if folder.name in ("hindi", "sanskrit") else "Arial"
        body_font = Font(name=font_name, size=10)
        ws = wb.create_sheet(subject[:31])
        ws.append(["Subject", "Topic", "Name", "Script", "Duration"])
        for cell in ws[1]:
            cell.font, cell.fill, cell.border = head_font, head_fill, border
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 24

        for t in topics:
            script, seconds = render_script(t)
            ws.append([
                subject,
                f"Ch {t['chapter_no']}: {t['chapter']}",
                f"{t['chapter_no']}.{t['topic_no']} {t['name']}",
                script,
                fmt(seconds),
            ])
            row = ws.max_row
            for cell in ws[row]:
                cell.font, cell.alignment, cell.border = body_font, top_wrap, border
            ws.cell(row, 5).alignment = Alignment(horizontal="center", vertical="top")
            # Excel won't auto-size rows written by openpyxl; estimate from wrapped lines (cap 409 pt).
            lines = sum(max(1, math.ceil(len(l) / 105)) for l in script.split("\n"))
            ws.row_dimensions[row].height = min(409, 13 * lines + 6)

        for col, width in zip("ABCDE", (14, 34, 40, 110, 11)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    out = ROOT / f"Kanthast_Class_{ROMAN.get(class_no, class_no)}.xlsx"
    wb.save(out)
    return out


if __name__ == "__main__":
    print(build(int(sys.argv[1])))
